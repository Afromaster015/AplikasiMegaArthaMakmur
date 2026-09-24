"""Modul laporan: laporan internal & konsultan + export Excel."""
import os

import payroll.db as db
import payroll.rekap as rekap

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None


def laporan_internal(db_path=None, id_periode=None):
    """Laporan internal lengkap (semua komponen per pegawai)."""
    return rekap.rekap_periode(db_path, id_periode)


def laporan_konsultan(db_path=None, id_periode=None):
    """Laporan ringkas untuk konsultan.

    Kolom: kode, nama, hari masuk, gaji bulanan, upah harian, uang makan,
    lembur, total tunjangan, potongan, total bersih (netto).
    """
    hasil = rekap.rekap_periode(db_path, id_periode)
    baris = []
    for b in hasil["baris"]:
        tunjangan = b["uang_makan"] + b["tambahan_tgl_merah"] + \
            b["uang_makan_lembur"] + b["lembur_malam"] + b["lembur_jam"]
        baris.append({
            "id_karyawan": b["id_karyawan"],
            "nama": b["nama"],
            "tipe_gaji": b["tipe_gaji"],
            "hari_masuk": b["hari_masuk"],
            "gaji": b["gaji_bulanan"],
            "upah": b["upah_harian"],
            "uang_makan": b["uang_makan"],
            "lembur": b["uang_makan_lembur"] + b["lembur_malam"] + b["lembur_jam"],
            "total_tunjangan": tunjangan,
            "potongan": b["potongan"],
            "total_bersih": b["total_bersih"],
        })
    return {"baris": baris, "total": hasil["total"]}


def export_excel(db_path=None, id_periode=None, tujuan=None, jenis="konsultan"):
    """Export laporan ke file Excel (.xlsx). Return path file."""
    if openpyxl is None:
        raise RuntimeError("openpyxl belum terinstall — jalankan: pip install openpyxl")

    conn = db.get_conn(db_path)
    try:
        p = conn.execute("SELECT * FROM tb_periode WHERE id_periode=?",
                         (id_periode,)).fetchone()
        nama_periode = p["bulan"] if p else str(id_periode)
    finally:
        conn.close()

    if jenis == "konsultan":
        data = laporan_konsultan(db_path, id_periode)
        header = ["No", "Kode", "Nama", "Hari Kerja", "Gaji Bulanan", "Upah Harian",
                  "Uang Makan", "Lembur", "Total Tunjangan", "Potongan",
                  "Total Gaji/Upah (Netto)"]
    else:
        data = laporan_internal(db_path, id_periode)
        header = ["No", "Kode", "Nama", "Hari", "Gaji", "Upah", "U. Makan",
                  "Tgl Merah", "UML", "ULM", "ULJ", "Potongan", "Total Bersih"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan"

    # Judul
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
    judul = ws.cell(row=1, column=1, value=f"PT. MEGA ARTHA MAKMUR — {nama_periode.upper()}")
    judul.font = Font(bold=True, size=13)
    judul.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(header))
    sub = ws.cell(row=2, column=1,
                  value="REKAP GAJI DAN UPAH (LAPORAN KONSULTAN)" if jenis == "konsultan"
                  else "REKAP GAJI DAN UPAH (LAPORAN INTERNAL)")
    sub.font = Font(bold=True, size=10, color="555555")
    sub.alignment = Alignment(horizontal="center")

    # Header
    for i, h in enumerate(header, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=Side(style="thin"))

    # Data
    thin = Border(bottom=Side(style="thin", color="DDDDDD"))
    for idx, b in enumerate(data["baris"], 1):
        r = idx + 4
        if jenis == "konsultan":
            vals = [idx, b["id_karyawan"], b["nama"], b["hari_masuk"], b["gaji"],
                    b["upah"], b["uang_makan"], b["lembur"], b["total_tunjangan"],
                    b["potongan"], b["total_bersih"]]
        else:
            vals = [idx, b["id_karyawan"], b["nama"], b["hari_masuk"], b["gaji_bulanan"],
                    b["upah_harian"], b["uang_makan"], b["tambahan_tgl_merah"],
                    b["uang_makan_lembur"], b["lembur_malam"], b["lembur_jam"],
                    b["potongan"], b["total_bersih"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            if i > 3:
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
            if i in (3,):
                pass
            c.border = thin

    # Total
    r = len(data["baris"]) + 5
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    if jenis == "konsultan":
        ws.cell(row=r, column=5, value=data["total"]["gaji"]).font = Font(bold=True)
        ws.cell(row=r, column=6, value=data["total"]["upah"]).font = Font(bold=True)
        ws.cell(row=r, column=7, value=data["total"]["uang_makan"]).font = Font(bold=True)
        ws.cell(row=r, column=8, value=data["total"]["lembur"]).font = Font(bold=True)
        ws.cell(row=r, column=10, value=data["total"]["potongan"]).font = Font(bold=True)
        ws.cell(row=r, column=11, value=data["total"]["total_bersih"]).font = Font(bold=True)

    # Lebar kolom
    lebar = [5, 10, 26, 10, 14, 14, 12, 12, 14, 12, 16]
    for i, w in enumerate(lebar[:len(header)], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(tujuan) or ".", exist_ok=True)
    wb.save(tujuan)
    return tujuan

